from __future__ import annotations

import hashlib
import io
import json
import os
import uuid
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

PACKAGE_ROOT = Path(__file__).resolve().parents[1] / "src" / "cnie_documents"
ROOT = PACKAGE_ROOT / "templates"
COAT_OF_ARMS = PACKAGE_ROOT / "assets" / "royaume-du-maroc.png"
AMIRI_QURAN_FONT = PACKAGE_ROOT / "assets" / "fonts" / "AmiriQuran-1.003.ttf"
AMIRI_QURAN_NAME = "Amiri Quran"
AMIRI_QURAN_SHA256 = "e2a47644762d16bdfb6d33e0d8db8c6ff30beae84150ef5a705316bbd829455c"
MARRIAGE_FIRST = PACKAGE_ROOT / "assets" / "background-marriage-first-v2.png"
MARRIAGE_CONTINUATION = PACKAGE_ROOT / "assets" / "background-marriage-continuation.png"
INHERITANCE_FIRST = PACKAGE_ROOT / "assets" / "background-inheritance-first.png"
INHERITANCE_CONTINUATION = PACKAGE_ROOT / "assets" / "background-inheritance-continuation.png"
ARABIC_FONT = "Arial"
MOROCCO_GREEN = "006233"
ANTIQUE_GOLD = "B59A5A"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)

MARRIAGE_VERSE = (
    "وَمِنْ ءَايَـٰتِهِۦٓ أَنْ خَلَقَ لَكُم مِّنْ أَنفُسِكُمْ أَزْوَٰجًۭا لِّتَسْكُنُوٓا۟ "
    "إِلَيْهَا وَجَعَلَ بَيْنَكُم مَّوَدَّةًۭ وَرَحْمَةً ۚ إِنَّ فِى ذَٰلِكَ "
    "لَـَٔايَـٰتٍۢ لِّقَوْمٍۢ يَتَفَكَّرُونَ"
)


def rtl_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.RIGHT, after=3, line=1.05,
                  bidi=True):
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    properties = paragraph._p.get_or_add_pPr()
    if bidi and properties.find(qn("w:bidi")) is None:
        properties.append(OxmlElement("w:bidi"))
    return paragraph


def style_run(run, size=13, bold=False, color="000000", rtl=True,
              font_name=ARABIC_FONT):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), font_name)
    complex_size = props.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = OxmlElement("w:szCs")
        props.append(complex_size)
    complex_size.set(qn("w:val"), str(int(size * 2)))
    if rtl and props.find(qn("w:rtl")) is None:
        props.append(OxmlElement("w:rtl"))
    return run


def add_text(paragraph, text: str, *, size=13, bold=False, color="000000", rtl=True,
             font_name=ARABIC_FONT):
    return style_run(paragraph.add_run(text), size=size, bold=bold, color=color,
                     rtl=rtl, font_name=font_name)


def add_control(paragraph, tag: str, alias: str, placeholder="................", *, size=13,
                bold=False, rtl=True, automatic_placeholder: str | None = None):
    control = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    alias_node = OxmlElement("w:alias")
    alias_node.set(qn("w:val"), alias)
    tag_node = OxmlElement("w:tag")
    tag_node.set(qn("w:val"), tag)
    identifier = OxmlElement("w:id")
    stable_id = int.from_bytes(hashlib.sha256(tag.encode("utf-8")).digest()[:4], "big")
    identifier.set(qn("w:val"), str(stable_id % 2_000_000_000 + 1))
    appearance = etree.Element(f"{{{W15}}}appearance")
    appearance.set(f"{{{W15}}}val", "hidden")
    props.extend((alias_node, tag_node, identifier, OxmlElement("w:text"), appearance))
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), ARABIC_FONT)
    run_props.append(fonts)
    if rtl:
        run_props.append(OxmlElement("w:rtl"))
    if bold:
        run_props.append(OxmlElement("w:b"))
        run_props.append(OxmlElement("w:bCs"))
    for node_name in ("sz", "szCs"):
        size_node = OxmlElement(f"w:{node_name}")
        size_node.set(qn("w:val"), str(int(size * 2)))
        run_props.append(size_node)
    run.append(run_props)
    text = OxmlElement("w:t")
    text.text = (automatic_placeholder if automatic_placeholder is not None else " ") \
        if tag.startswith("enotario.") else placeholder
    if text.text.startswith(" ") or text.text.endswith(" "):
        text.set(qn("xml:space"), "preserve")
    run.append(text)
    content.append(run)
    control.extend((props, content))
    paragraph._p.append(control)


def configure_document(document: Document, *, marriage=False) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.8 if marriage else 2.55)
    # Leave the printed footer ornament outside the editable text area.
    section.bottom_margin = Cm(2.8 if marriage else 4.5)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.35)
    section.footer_distance = Cm(0.85)
    section.different_first_page_header_footer = True
    normal = document.styles["Normal"]
    normal.font.name = ARABIC_FONT
    normal.font.size = Pt(13)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""


def _paragraph_border(paragraph, *, color: str, size=6, space=4,
                      edges=("top", "bottom")) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    for edge in edges:
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), str(size))
        item.set(qn("w:space"), str(space))
        item.set(qn("w:color"), color)
        borders.append(item)


def _paragraph_fill(paragraph, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _anchor_inline(inline, *, behind=True) -> None:
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "1" if behind else "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)
    for axis in ("H", "V"):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        alignment = OxmlElement("wp:align")
        alignment.text = "center"
        position.append(alignment)
        anchor.append(position)
    for child in list(inline):
        anchor.append(deepcopy(child))
    wrap = OxmlElement("wp:wrapNone")
    doc_property = anchor.find(qn("wp:docPr"))
    insertion = list(anchor).index(doc_property) if doc_property is not None else len(anchor)
    anchor.insert(insertion, wrap)
    inline.getparent().replace(inline, anchor)


def _add_anchored_picture(paragraph, path: Path, *, width, height=None,
                          opacity: int | None = None, doc_id: int) -> None:
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=width, height=height)
    inline = shape._inline
    inline.docPr.set("id", str(doc_id))
    if opacity is not None:
        blips = inline.xpath(".//a:blip")
        if blips:
            alpha = OxmlElement("a:alphaModFix")
            alpha.set("amt", str(opacity))
            blips[0].append(alpha)
    _anchor_inline(inline)


def _prepare_header(header, frame: Path, *, watermark=False, compact=False,
                    verse_text: str = "", frame_id: int) -> None:
    paragraph = header.paragraphs[0]
    paragraph.clear()
    rtl_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0, bidi=False)
    paragraph.paragraph_format.space_before = Pt(0)
    # A centered 5 mm safety inset protects the supplied A4 border on printers
    # that cannot print to the edge. Its aspect ratio remains A4.
    _add_anchored_picture(paragraph, frame, width=Cm(20), height=Cm(28.28),
                          doc_id=frame_id)
    # A second anchored drawing keeps both first/continuation page backgrounds
    # active in Word. It is visibly used only for the marriage first page.
    _add_anchored_picture(paragraph, COAT_OF_ARMS, width=Cm(8.2),
                          opacity=7000 if watermark else 2000,
                          doc_id=frame_id + 1)
    if verse_text:
        verse = rtl_paragraph(header.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER,
                              after=0, line=1.06)
        verse.paragraph_format.space_before = Cm(1.85)
        verse.paragraph_format.left_indent = Cm(1.65)
        verse.paragraph_format.right_indent = Cm(1.65)
        run = add_text(verse, verse_text, size=10.5, color=MOROCCO_GREEN,
                       font_name=AMIRI_QURAN_NAME)
        language = OxmlElement("w:lang")
        language.set(qn("w:val"), "ar-SA")
        run._element.get_or_add_rPr().append(language)
    if compact:
        # A real continuation line keeps the Word header active. Unlike the
        # former 3.5 mm placement, it sits safely below the printed top frame.
        line = rtl_paragraph(header.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER,
                             after=0, line=1.0)
        run = add_text(line, "المملكة المغربية  ·  وزارة العدل  ·  قسم التوثيق", size=8.5,
                       bold=True, color=MOROCCO_GREEN)
        baseline = OxmlElement("w:position")
        baseline.set(qn("w:val"), "-40")
        run._element.get_or_add_rPr().append(baseline)


def _prepare_footer(footer, *, marriage=False) -> None:
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    rtl_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0,
                  bidi=not marriage)
    _paragraph_border(paragraph, color=ANTIQUE_GOLD if marriage else MOROCCO_GREEN,
                      size=5, space=4, edges=("top",))
    if marriage:
        add_text(paragraph, "ROYAUME DU MAROC\nMINISTÈRE DE LA JUSTICE", size=7.5,
                 bold=True, color=MOROCCO_GREEN, rtl=False)
    else:
        add_text(paragraph, "المملكة المغربية\nالعدالة في خدمة المواطن", size=8,
                 bold=True, color=MOROCCO_GREEN)


def add_page_identity(document: Document, first: Path, continuation: Path,
                      *, verse_text: str = "") -> None:
    section = document.sections[0]
    _prepare_header(section.first_page_header, first, verse_text=verse_text,
                    frame_id=101)
    _prepare_header(section.header, continuation, compact=True, frame_id=201)
    # The supplied continuation artwork already carries the visual footer.
    # A second Word footer would overlap it on dense continuation pages.


def add_first_page_spacer(document: Document, height: float,
                          *, hidden_legal_text: str = "") -> None:
    """Reserve the p1 artwork's header/hero without repeating it in Word text."""
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Cm(height)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    if hidden_legal_text:
        # Keep the p1 verse searchable for the immutable legal-text check,
        # without printing a duplicate over the supplied hero artwork.
        run = add_text(spacer, hidden_legal_text, size=1)
        run._element.get_or_add_rPr().append(OxmlElement("w:vanish"))
    spacer.add_run(" ").font.size = Pt(1)


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)


def _set_cell_margins(cell, value=40) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)


def add_official_header(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _remove_table_borders(table)
    widths = (Cm(6.35), Cm(3.4), Cm(6.35))
    cells = table.rows[0].cells
    for cell, width in zip(cells, widths, strict=True):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell)
        paragraph = rtl_paragraph(cell.paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)

    add_text(cells[0].paragraphs[0], "المحكمة الابتدائية", size=14.5, bold=True,
             color=MOROCCO_GREEN)
    add_text(cells[0].paragraphs[0], "\n◆\n", size=7.5, color=ANTIQUE_GOLD)
    add_text(cells[0].paragraphs[0], "قسم التوثيق", size=11.5, bold=True)

    logo_run = cells[1].paragraphs[0].add_run()
    shape = logo_run.add_picture(str(COAT_OF_ARMS), width=Cm(2.75))
    shape._inline.docPr.set("id", "301")
    shape._inline.docPr.set("descr", "شعار المملكة المغربية")
    shape._inline.docPr.set("title", "المملكة المغربية")

    add_text(cells[2].paragraphs[0], "المملكة المغربية", size=14.5, bold=True,
             color=MOROCCO_GREEN)
    add_text(cells[2].paragraphs[0], "\n◆\n", size=7.5, color=ANTIQUE_GOLD)
    add_text(cells[2].paragraphs[0], "وزارة العدل", size=11.5, bold=True)


def add_title(document: Document, text: str) -> None:
    paragraph = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.left_indent = Cm(5.2)
    paragraph.paragraph_format.right_indent = Cm(5.2)
    _paragraph_border(paragraph, color=ANTIQUE_GOLD, size=7, space=5)
    add_text(paragraph, "◇  ", size=11, bold=True, color=ANTIQUE_GOLD)
    add_text(paragraph, text, size=38, bold=True, color=MOROCCO_GREEN)
    add_text(paragraph, "  ◇", size=11, bold=True, color=ANTIQUE_GOLD)


def add_registry_line(document: Document, book: str) -> None:
    paragraph = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.0)
    add_text(paragraph, "ضمن بعدد ", size=11.5, bold=True)
    add_control(paragraph, "manual.registry_number", "رقم التضمين", size=11.5)
    add_text(paragraph, "  صحيفة ", size=11.5, bold=True)
    add_control(paragraph, "manual.page", "الصحيفة", size=11.5)
    add_text(paragraph, f"  {book} رقم ", size=11.5, bold=True)
    add_control(paragraph, "manual.book_number", "رقم الكناش", size=11.5)
    add_text(paragraph, "  بتاريخ ", size=11.5, bold=True)
    add_control(paragraph, "manual.registry_date", "تاريخ التضمين", size=11.5)

    divider = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER,
                            after=6, line=1.0, bidi=False)
    add_text(divider, "────────────  ◇  ────────────", size=9, bold=True,
             color=ANTIQUE_GOLD, rtl=False)


def build_marriage() -> Document:
    document = Document()
    configure_document(document, marriage=True)
    add_page_identity(document, MARRIAGE_FIRST, MARRIAGE_CONTINUATION,
                      verse_text=MARRIAGE_VERSE)
    add_first_page_spacer(document, 7.8, hidden_legal_text=MARRIAGE_VERSE)
    add_registry_line(document, "كناش الزواج")

    body = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=5, line=1.2)
    size = 12.5
    add_text(body, " الحمد لله وحده وصل عدلي ", size=size)
    add_control(body, "manual.notary_1", "العدل الأول", size=size)
    add_text(body, "تلقى العدلان ", size=size)
    add_control(body, "manual.notary_2", "العدل الثاني", size=size)
    add_text(body, "و", size=size)
    add_control(body, "manual.notary_3", "العدل المشارك", size=size)
    add_text(body, "المنتصبان للإشهاد بالمحكمة الابتدائية بالجديدة الإشهاد الآتي المدرج نصــــــــــه ", size=size)
    add_text(body, "بناء على إذن بتوثيق عقد الزواج", size=size, bold=True)
    add_text(body, " من السيد قاضي الأسرة المكلف بالزواج بقسم قضاء الأسرة بالمحكمة الابتدائية بالجديدة ملف رقم ", size=size)
    add_control(body, "manual.case_number", "رقم الملف", size=size)
    add_text(body, " بتاريخ ", size=size)
    add_control(body, "manual.case_date", "تاريخ الملف", size=size)
    add_text(body, " تـــــــزوج على بركة الله وحسن عونه الجميل السيــــــــــد ", size=size, bold=True)
    add_control(body, "enotario.husband.1.person_ar", "بيانات الزوج", size=size)
    add_text(body, "، البنــــــــــت ", size=size, bold=True)
    add_control(body, "enotario.wife.1.person_ar", "بيانات الزوجة", size=size)
    add_text(body, "، الحل للزواج الخالية من موانعه الشرعية على صداق مبارك قدره ", size=size)
    add_control(body, "manual.dowry", "الصداق", size=size)
    add_text(body, " اعترفت الزوجة ووالدها بقبض جميع الصداق المذكور و أبرآه منه أتم إبراء تزوجها وفق الكتاب والسنة واليمن والأمان وما جاء في محكم القرآن عقد زواجها والدها السيــــــــــد ", size=size)
    add_control(body, "enotario.wife_father.1.person_ar", "بيانات والد الزوجة", size=size)
    add_text(body, "، بتفويض منها له على ذلك سمعه منها شهيداه في تاريخه بعد صدور الإيجاب والقبول من المتعاقدين وهما متمتعان بالأهلية والتمييز والاختيار وأشعر الطرفان من طرف شهيديه بأنه يجوز لهما في إطار تدبير الأموال التي ستكتسب أثناء قيام الزوجية الاتفاق على استثمارها وتوزيعها في وثيقة مستقلة عن عقد الزواج والله سبحانه يؤلف بينهما لما يحبه ويرضاه وتلي عليهما مضمون هذا العقد عرفوا قدره وهم بأتمه وعرف بهم وحرر بتاريخه", size=size)
    return document


def build_inheritance() -> Document:
    document = Document()
    configure_document(document)
    add_page_identity(document, INHERITANCE_FIRST, INHERITANCE_CONTINUATION)
    add_first_page_spacer(document, 4.15)
    add_registry_line(document, "كناش التركـــــات")

    body = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, line=1.18)
    size = 12
    add_text(body, "الحمد لله وحده وصل عدلي ", size=size)
    add_control(body, "manual.notary_1", "العدل الأول", size=size)
    add_text(body, "تلقى العدلان ", size=size)
    add_control(body, "manual.notary_2", "العدل الثاني", size=size)
    add_text(body, "و", size=size)
    add_control(body, "manual.notary_3", "العدل المشارك", size=size)
    add_text(body, "المنتصبان للإشهاد بالمحكمة الابتدائية بالجديدة الإشهاد الآتي المدرج نصــــــــــه بطلب من ", size=size)
    add_control(body, "enotario.applicant.1.person_ar", "بيانات طالب الشهادة", size=size,
                automatic_placeholder="السيــــــد ")
    add_text(body, " شهوده الآتي ذكرهم ذكروا أنهم يعرفون المرحوم بالله السيـــــــد ", size=size)
    add_control(body, "manual.deceased", "اسم الهالك", size=size)
    add_text(body, "الذي كان ساكنا قيد حياته بنفس العنوان المذكور أعلاه المعرفة التامة الكافية شرعا بها ومعها يشهدون بأنه توفي رحمه الله بتاريخ ", size=size)
    add_control(body, "manual.death_date", "تاريخ الوفاة", size=size)
    add_text(body, "حسب نسخة موجزة من رسم الوفاة رقم ", size=size)
    add_control(body, "manual.death_record", "رسم الوفاة", size=size)
    add_text(body, "بتاريخ ", size=size)
    add_control(body, "manual.death_record_date", "تاريخ رسم الوفاة", size=size)
    add_text(body, "فأحاط بإرثه ", size=size, bold=True)
    for index in range(1, 13):
        add_control(body, f"enotario.heir.{index}.person_ar", f"بيانات الوارث {index}", size=size)
        add_control(body, f"manual.heir_relation.{index}", f"صلة الوارث {index}",
                    placeholder=" ", size=size)
    add_text(body, "، ولا يعلمون له وارثا سوى من ذكر كما لا يعلمون له وصية ولا تنزيلا ولا ابن  ابن أو ابن بنت توفيا قيد حياته ويعرفون الورثة المذكورين مثل معرفة موروثهم المذكور هذا ما في علمهم وصحة يقينهم علموه بعضهم بالقرابة والبعض بالمجاورة والكل بالمخالطة وشدة الاطلاع على الأحوال وبمضمنه قيدت شهادتهم مسؤولة منهم لسائلها أعلاه وهم عارفون قدره وبأتمه وعرف بهـــم وبمقتضــى ما ذكر فقد صحت فريضة الهالك المذكور من ", size=size)
    add_control(body, "manual.estate_basis", "أصل الفريضة", size=size)
    add_text(body, "شهد بذلك السادة ", size=size, bold=True)

    for index in range(1, 13):
        # Word mirrors left/right alignment for bidi paragraphs. LEFT therefore
        # anchors this RTL witness line to the visual right edge, as in the source.
        witness = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.LEFT,
                                after=2, line=1.08)
        witness.paragraph_format.keep_together = True
        add_text(witness, f"-{index}- ", size=11.5, bold=True, color=MOROCCO_GREEN)
        add_control(witness, f"enotario.witness.{index}.person_ar", f"بيانات الشاهد {index}", size=11.5)

    closing = rtl_paragraph(document.add_paragraph(), align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, line=1.03)
    closing.paragraph_format.space_before = Pt(3)
    add_text(closing, "والشهود المذكورون كلهم مغاربة جنسية وتلي مضمون الشهادة على طالب الشهادة أعلاه والشهــــود المذكوريـن وحرر بتاريخـــه عبـــــد ربـــه تعالـــــى                 وعبـــد ربه تعالـــــى ", size=size)
    return document


def normalize_package(path: Path) -> None:
    """Remove creation-time entropy so a template version has one stable hash."""
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    core_name = "docProps/core.xml"
    if core_name in parts:
        root = etree.fromstring(parts[core_name])
        ns = {"dcterms": "http://purl.org/dc/terms/"}
        for node in root.xpath("./dcterms:created | ./dcterms:modified", namespaces=ns):
            node.text = "2026-01-01T00:00:00Z"
        parts[core_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as target:
        for name in sorted(parts):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, parts[name])
    path.write_bytes(output.getvalue())


def embed_verse_font(path: Path) -> None:
    """Bundle the OFL-licensed Quran font in Word's native obfuscated font part."""
    font_payload = AMIRI_QURAN_FONT.read_bytes()
    if hashlib.sha256(font_payload).hexdigest() != AMIRI_QURAN_SHA256:
        raise ValueError("Amiri Quran font integrity check failed")
    key = uuid.UUID(bytes=hashlib.sha256(font_payload).digest()[:16])
    reversed_key = key.bytes[::-1]
    obfuscated = bytearray(font_payload)
    for index in range(32):
        obfuscated[index] ^= reversed_key[index % 16]
    font_part = "word/fonts/AmiriQuran.odttf"
    relationship_id = "rIdAmiriQuran"
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}

    font_table = etree.fromstring(parts["word/fontTable.xml"])
    existing = font_table.xpath("./w:font[@w:name=$name]", namespaces={"w": W},
                                name=AMIRI_QURAN_NAME)
    font_node = existing[0] if existing else etree.SubElement(
        font_table, f"{{{W}}}font", {f"{{{W}}}name": AMIRI_QURAN_NAME})
    etree.SubElement(font_node, f"{{{W}}}embedRegular", {
        f"{{{R}}}id": relationship_id,
        f"{{{W}}}fontKey": "{" + str(key).upper() + "}",
    })
    parts["word/fontTable.xml"] = etree.tostring(
        font_table, xml_declaration=True, encoding="UTF-8", standalone=True)

    rels_name = "word/_rels/fontTable.xml.rels"
    relationships = etree.fromstring(parts[rels_name]) if rels_name in parts else etree.Element(
        f"{{{PKG_REL}}}Relationships", nsmap={None: PKG_REL})
    if any(item.get("Id") == relationship_id for item in relationships):
        raise ValueError("Embedded font relationship is duplicated")
    etree.SubElement(relationships, f"{{{PKG_REL}}}Relationship", {
        "Id": relationship_id,
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
        "Target": "fonts/AmiriQuran.odttf",
    })
    parts[rels_name] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True)

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(content_types, f"{{{CT}}}Override", {
        "PartName": "/" + font_part,
        "ContentType": "application/vnd.openxmlformats-officedocument.obfuscatedFont",
    })
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True)
    settings = etree.fromstring(parts["word/settings.xml"])
    if settings.find(f"{{{W}}}embedTrueTypeFonts") is None:
        etree.SubElement(settings, f"{{{W}}}embedTrueTypeFonts", {f"{{{W}}}val": "1"})
    parts["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone=True)
    parts[font_part] = bytes(obfuscated)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for name, payload in parts.items():
            package.writestr(name, payload)


def main() -> None:
    output_root = Path(os.environ.get("ENOTARIO_TEMPLATE_OUTPUT_ROOT", ROOT))
    builders = {"marriage": build_marriage, "inheritance": build_inheritance}
    selected = tuple(filter(None, os.environ.get("ENOTARIO_TEMPLATE_NAMES", "").split(",")))
    names = selected or tuple(builders)
    if set(names) - set(builders):
        raise ValueError("Unknown template name")
    for name in names:
        builder = builders[name]
        source_folder = ROOT / name
        folder = output_root / name
        folder.mkdir(parents=True, exist_ok=True)
        destination = folder / "template.docx"
        document = builder()
        document.save(destination)
        if name == "marriage":
            embed_verse_font(destination)
        normalize_package(destination)
        print(json.dumps({
            "template": str(destination),
            "template_sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
            "legal_text_sha256": hashlib.sha256((source_folder / "legal.txt").read_bytes()).hexdigest(),
        }, ensure_ascii=False))


if __name__ == "__main__":
    main()
